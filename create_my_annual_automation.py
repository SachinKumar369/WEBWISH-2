from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
DEEP_BLUE = RGBColor(0x0D, 0x2C, 0x54)
MED_BLUE = RGBColor(0x1A, 0x56, 0x9E)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF0, 0xF4, 0xF8)
GREEN = RGBColor(0x27, 0xAE, 0x60)
GREEN_LIGHT = RGBColor(0xE8, 0xF8, 0xEF)
DARK_TEXT = RGBColor(0x1C, 0x1C, 0x1C)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GOLD = RGBColor(0xD4, 0xA0, 0x1E)

def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, w, h, color, shape_type=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape_type, left, top, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s

def add_text(slide, left, top, w, h, text, size=18, bold=False, color=DARK_TEXT, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tb

# =========================================================
# SLIDE 1 — TITLE & SUMMARY
# =========================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide1, DEEP_BLUE)

# Top accent bar
add_shape(slide1, Inches(0), Inches(0), Inches(13.333), Inches(0.08), ACCENT_BLUE)

# Title
add_text(slide1, Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.2),
         "One Year of QA Automation Achievements", size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# Subtitle
add_text(slide1, Inches(1.5), Inches(1.6), Inches(10.3), Inches(0.7),
         "Automated Testing  |  Team Enablement  |  Faster Delivery", size=22, color=ACCENT_BLUE, align=PP_ALIGN.CENTER)

# Divider line
add_shape(slide1, Inches(4.5), Inches(2.4), Inches(4.3), Inches(0.04), GOLD)

# ---- 4 Key Cards ----
cards = [
    ("WebProlific", "Inventory & Accounting\nModule Automation", ACCENT_BLUE),
    ("Flutter App", "Mobile E2E Test\nAutomation", MED_BLUE),
    ("WebWish PMS", "Full-Stack Front Desk,\nMarketing & Reports", DEEP_BLUE),
    ("Team Training", "Upskilled Team on\nPlaywright & Automation", GREEN),
]
card_w = Inches(2.7)
card_h = Inches(2.2)
gap = Inches(0.35)
start_x = Inches(0.85)
card_y = Inches(3.0)

for i, (title, desc, clr) in enumerate(cards):
    x = start_x + i * (card_w + gap)
    c = add_shape(slide1, x, card_y, card_w, card_h, clr, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide1, x, card_y + Inches(0.3), card_w, Inches(0.5), title, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide1, x + Inches(0.15), card_y + Inches(0.95), card_w - Inches(0.3), Inches(1.0), desc, size=15, color=WHITE, align=PP_ALIGN.CENTER)

# Bottom impact bar
add_shape(slide1, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.1), RGBColor(0x0A, 0x22, 0x42))
add_text(slide1, Inches(0.8), Inches(5.65), Inches(11.7), Inches(0.5),
         "KEY IMPACT", size=14, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

# Two impact stats
add_text(slide1, Inches(0.8), Inches(6.05), Inches(5.85), Inches(0.6),
         "3  Applications Automated  (WebProlific + Flutter + WebWish)", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide1, Inches(6.65), Inches(6.05), Inches(5.85), Inches(0.6),
         "Significant Reduction in Regression Cycle Time", size=16, bold=True, color=GREEN, align=PP_ALIGN.CENTER)


# =========================================================
# SLIDE 2 — DETAILS & RESULTS
# =========================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2, LIGHT_BG)

# Header bar
add_shape(slide2, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DEEP_BLUE)
add_text(slide2, Inches(0.8), Inches(0.18), Inches(11.5), Inches(0.7),
         "Automation Work Summary  —  Year in Review", size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ---- Left Column: What Was Automated ----
col1_x = Inches(0.6)
col1_w = Inches(5.9)

add_shape(slide2, col1_x, Inches(1.3), col1_w, Inches(5.7), WHITE)
add_text(slide2, col1_x + Inches(0.3), Inches(1.45), col1_w - Inches(0.6), Inches(0.5),
         "What Was Automated", size=20, bold=True, color=DEEP_BLUE)

items_left = [
    ("WebProlific — Inventory Module", "Automated inventory workflows covering stock management, purchase orders, and stock adjustments using Playwright."),
    ("WebProlific — Accounting Module", "Automated accounting flows including journal entries, ledgers, and financial report generation."),
    ("Flutter Mobile App", "Built E2E automation for the Flutter mobile application covering key user journeys."),
    ("WebWish PMS", "Automated Front Desk, Marketing, Guest Management, Reports, Room Inventory, and Global Search modules."),
]

y_pos = Inches(2.1)
for title, desc in items_left:
    # Bullet icon
    add_shape(slide2, col1_x + Inches(0.35), y_pos + Inches(0.05), Inches(0.14), Inches(0.14), ACCENT_BLUE, MSO_SHAPE.OVAL)
    add_text(slide2, col1_x + Inches(0.6), y_pos - Inches(0.05), col1_w - Inches(1.0), Inches(0.35),
             title, size=14, bold=True, color=DARK_TEXT)
    add_text(slide2, col1_x + Inches(0.6), y_pos + Inches(0.28), col1_w - Inches(1.0), Inches(0.55),
             desc, size=12, color=RGBColor(0x55, 0x55, 0x55))
    y_pos += Inches(0.95)

# ---- Right Column: Team Enablement & Impact ----
col2_x = Inches(6.8)
col2_w = Inches(5.9)

# Team Training Card
add_shape(slide2, col2_x, Inches(1.3), col2_w, Inches(2.6), WHITE)
add_text(slide2, col2_x + Inches(0.3), Inches(1.45), col2_w - Inches(0.6), Inches(0.5),
         "Team Enablement", size=20, bold=True, color=DEEP_BLUE)

training_items = [
    "Conducted hands-on training sessions for the QA team",
    "Covered Playwright framework, test design & best practices",
    "Enabled team to independently write & maintain automation tests",
    "Established shared test patterns and reusable utilities",
]
ty = Inches(2.1)
for t in training_items:
    add_text(slide2, col2_x + Inches(0.35), ty, col2_w - Inches(0.7), Inches(0.4),
             "✓  " + t, size=13, color=DARK_TEXT)
    ty += Inches(0.38)

# Impact Card
add_shape(slide2, col2_x, Inches(4.15), col2_w, Inches(2.85), GREEN_LIGHT)
add_text(slide2, col2_x + Inches(0.3), Inches(4.3), col2_w - Inches(0.6), Inches(0.5),
         "Measurable Impact", size=20, bold=True, color=GREEN)

impact_items = [
    ("Regression Time", "Reduced from days to hours"),
    ("Defect Catch Rate", "Earlier bug detection in CI pipeline"),
    ("Release Confidence", "Faster, reliable release sign-offs"),
    ("Team Velocity", "Team now self-sufficient in automation"),
]
iy = Inches(4.95)
for label, val in impact_items:
    add_shape(slide2, col2_x + Inches(0.35), iy + Inches(0.05), Inches(0.14), Inches(0.14), GREEN, MSO_SHAPE.OVAL)
    add_text(slide2, col2_x + Inches(0.6), iy - Inches(0.05), Inches(2.2), Inches(0.35),
             label, size=13, bold=True, color=DEEP_BLUE)
    add_text(slide2, col2_x + Inches(2.8), iy - Inches(0.05), Inches(2.8), Inches(0.35),
             val, size=13, color=DARK_TEXT)
    iy += Inches(0.45)

# Save
output_path = r"e:\Automation Project\WebWish 2\Annual_Automation_Achievements.pptx"
prs.save(output_path)
print(f"PPT saved to: {output_path}")
