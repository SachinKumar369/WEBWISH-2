"""
Generate ATS-friendly Word Resume for Avinash Tiwary
"""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

doc = Document()

# ─── Page Setup ───
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# ─── Style Definitions ───
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.space_before = Pt(0)

# Color palette
DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MEDIUM_BLUE = RGBColor(0x2C, 0x5F, 0x8A)
ACCENT_BLUE = RGBColor(0x3A, 0x7C, 0xBD)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)


def rgb_to_hex(color):
    """Convert RGBColor to hex string without '#'."""
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def add_horizontal_line(doc, color=ACCENT_BLUE):
    """Add a thin horizontal line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="{rgb_to_hex(color)}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def add_section_heading(doc, text):
    """Add a section heading with bottom border."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE
    # Add bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="8" w:space="1" w:color="{rgb_to_hex(ACCENT_BLUE)}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def add_subsection_heading(doc, left_text, right_text=""):
    """Add a subsection heading (e.g., Job Title + Company)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(left_text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = MEDIUM_BLUE
    if right_text:
        run2 = p.add_run(f"  |  {right_text}")
        run2.font.name = 'Calibri'
        run2.font.size = Pt(10)
        run2.font.color.rgb = LIGHT_GRAY


def add_project_heading(doc, title, role_tools):
    """Add a project heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = MEDIUM_BLUE
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    run2 = p2.add_run(role_tools)
    run2.font.name = 'Calibri'
    run2.font.size = Pt(9.5)
    run2.font.italic = True
    run2.font.color.rgb = MED_GRAY


def add_bullet(doc, text, bold_prefix="", indent_level=0):
    """Add a bullet point with optional bold prefix."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.8 + indent_level * 0.5)
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.font.name = 'Calibri'
        run_b.font.size = Pt(10.5)
        run_b.font.bold = True
        run_b.font.color.rgb = DARK_GRAY
        run_rest = p.add_run(text)
        run_rest.font.name = 'Calibri'
        run_rest.font.size = Pt(10.5)
        run_rest.font.color.rgb = DARK_GRAY
    else:
        # Clear default text and add with formatting
        p.clear()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_GRAY


def add_body_text(doc, text, bold=False, italic=False, size=Pt(10.5)):
    """Add a paragraph of body text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = DARK_GRAY
    return p


# ══════════════════════════════════════════════════════════════
#  HEADER — NAME & TITLE
# ══════════════════════════════════════════════════════════════

# Name
p_name = doc.add_paragraph()
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_name.paragraph_format.space_after = Pt(2)
run_name = p_name.add_run("AVINASH TIWARY")
run_name.font.name = 'Calibri'
run_name.font.size = Pt(22)
run_name.font.bold = True
run_name.font.color.rgb = DARK_BLUE

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(4)
run_title = p_title.add_run("Functional Test Analyst")
run_title.font.name = 'Calibri'
run_title.font.size = Pt(12)
run_title.font.color.rgb = ACCENT_BLUE

# Contact line
p_contact = doc.add_paragraph()
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_contact.paragraph_format.space_before = Pt(2)
p_contact.paragraph_format.space_after = Pt(4)
run_phone = p_contact.add_run("+91-9039591636")
run_phone.font.name = 'Calibri'
run_phone.font.size = Pt(10)
run_phone.font.color.rgb = MED_GRAY
run_sep = p_contact.add_run("  |  ")
run_sep.font.name = 'Calibri'
run_sep.font.size = Pt(10)
run_sep.font.color.rgb = LIGHT_GRAY
run_email = p_contact.add_run("iamavinashtiwary@outlook.com")
run_email.font.name = 'Calibri'
run_email.font.size = Pt(10)
run_email.font.color.rgb = MED_GRAY

add_horizontal_line(doc)

# ══════════════════════════════════════════════════════════════
#  PROFESSIONAL SUMMARY
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Professional Summary")

summary_text = (
    "Results-driven Functional Test Analyst with 5+ years of experience in software quality assurance "
    "across hotel hospitality systems, cloud-based, mobile, and Windows applications. Proven track record "
    "of reducing defects by 60% and increasing software stability by 40% through comprehensive test "
    "strategies. Skilled in manual testing, API testing (Postman), SQL database validation, and building/"
    "maintaining test automation frameworks with Playwright. Hands-on experience leveraging Agentic AI "
    "(GitHub Copilot, MCP Server, AI-assisted code generation) to accelerate test development, automate "
    "repetitive workflows, and improve overall QA efficiency. Recognized for self-learning and implementing "
    "new testing tools that improved team productivity by 25%."
)
add_body_text(doc, summary_text)

# ══════════════════════════════════════════════════════════════
#  CORE COMPETENCIES
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Core Competencies")

competencies = [
    "Manual Testing (Black Box, Regression, Integration, System, Stress, Ad-hoc, Exploratory)",
    "API Testing & Web Services (Postman)",
    "Test Automation (Playwright, Selenium)",
    "Test Framework Development & Maintenance (Page Object Model, Fixtures, Custom Reporters)",
    "Agentic AI-Driven Test Automation (GitHub Copilot, Playwright MCP Server, AI Code Generation)",
    "SQL Database Testing & Validation (SQL Server)",
    "Performance Testing (Apache JMeter)",
    "Test Planning, Test Case Design & Execution",
    "Defect Tracking & Lifecycle Management (TFS / Jira)",
    "Agile / Scrum Methodology | SDLC & STLC Lifecycle Management",
    "Test Data Preparation & Management",
    "Peer Review (Test Design & Defect Reporting)",
    "Exploratory Testing & Root Cause Analysis",
]

for comp in competencies:
    add_bullet(doc, comp)

# ══════════════════════════════════════════════════════════════
#  PROFESSIONAL EXPERIENCE
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Professional Experience")

add_subsection_heading(doc,
    "Functional Test Analyst",
    "Prologic First India Pvt. Ltd. | Gurugram, India | April 2018 – Present"
)

experience_bullets = [
    ("Designed and executed ", "comprehensive test plans for a cloud-based Hotel Property Management System (PMS), resulting in a 60% reduction in production defects and 40% improvement in software stability."),
    ("Developed and maintained ", "1,000+ manual test cases covering front desk, reporting, POS, and third-party integrations across multiple product lines."),
    ("Led ", "functional, integration, system, regression, API, and stress testing efforts across 3 major product lines (MyCloud PMS, Touche POS, and supporting modules)."),
    ("Performed API testing ", "using Postman for RESTful web services, improving backend reliability and reducing API-related defects by 35%."),
    ("Collaborated ", "with third-party developers (OPERA PMS, Shawman Table Reservation, Urban Piper) to conduct integration testing and ensure seamless interoperability."),
    ("Provided real-time support ", "to the support team for production issue resolution and POS application deployment."),
    ("Self-taught Postman ", "for API testing and trained the entire QA team, earning recognition from the Project Manager."),
    ("Implemented ", "structured testing principles, reducing testing cycle time by 50% and increasing team productivity by 25%."),
    ("Raised defects early ", "in the SDLC, enabling faster resolution and earning appreciation from the Project Manager."),
    ("Provided guidance ", "and supervision to junior testers, establishing standardized testing practices across the team."),
]

for bold_part, rest_part in experience_bullets:
    add_bullet(doc, rest_part, bold_prefix=bold_part)

# ── Sub-section: Framework Development ──
p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_before = Pt(8)
p_sub.paragraph_format.space_after = Pt(4)
run_sub = p_sub.add_run("Test Framework Development & Maintenance")
run_sub.font.name = 'Calibri'
run_sub.font.size = Pt(10.5)
run_sub.font.bold = True
run_sub.font.color.rgb = MEDIUM_BLUE
run_sub.font.italic = True

framework_bullets = [
    ("Designed and developed a scalable Playwright test automation framework ", "using TypeScript, Page Object Model (POM) pattern, and modular architecture for the WebWish PMS application."),
    ("Built and maintained ", "custom test fixtures, reusable utility libraries, and shared page objects to reduce code duplication and improve test maintainability."),
    ("Integrated Allure reporting, HTML reporters, and custom test dashboards ", "for real-time test execution visibility and stakeholder reporting."),
    ("Implemented CI/CD pipeline integration ", "(Jenkins/Docker) for automated test execution on code commits and release builds."),
    ("Established test data management strategies ", "including Excel-based data-driven testing, database seeding, and environment configuration management."),
    ("Refactored legacy test scripts ", "into a modular, maintainable framework reducing test maintenance effort by 40%."),
    ("Documented framework architecture, coding standards, and best practices ", "for team-wide adoption and onboarding."),
]

for bold_part, rest_part in framework_bullets:
    add_bullet(doc, rest_part, bold_prefix=bold_part)

# ── Sub-section: Agentic AI ──
p_sub2 = doc.add_paragraph()
p_sub2.paragraph_format.space_before = Pt(8)
p_sub2.paragraph_format.space_after = Pt(4)
run_sub2 = p_sub2.add_run("Agentic AI & Intelligent Test Automation")
run_sub2.font.name = 'Calibri'
run_sub2.font.size = Pt(10.5)
run_sub2.font.bold = True
run_sub2.font.color.rgb = MEDIUM_BLUE
run_sub2.font.italic = True

ai_bullets = [
    ("Leveraged GitHub Copilot ", "for AI-assisted test script generation, reducing test authoring time by 40% and accelerating framework development."),
    ("Integrated Playwright MCP (Model Context Protocol) Server ", "to enable AI-driven browser automation and intelligent test scenario generation."),
    ("Applied Agentic AI concepts ", "to automate repetitive QA workflows — including test case generation from requirements, automated regression suite creation, and intelligent defect categorization."),
    ("Used AI-assisted code review and refactoring ", "to improve test code quality, reduce technical debt, and ensure adherence to coding standards."),
    ("Developed AI-powered test data generators ", "and intelligent assertion builders that adapt to application changes with minimal manual intervention."),
    ("Explored natural language to test case translation ", "using LLM-based tools, enabling non-technical stakeholders to contribute to test planning."),
]

for bold_part, rest_part in ai_bullets:
    add_bullet(doc, rest_part, bold_prefix=bold_part)

# ══════════════════════════════════════════════════════════════
#  PROJECTS
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Projects")

# Project 1
add_project_heading(doc,
    "MyCloud Hospitality PMS System",
    "Role: Functional Test Engineer  |  Tools: SQL Server, TFS, Postman"
)

project1_bullets = [
    "Performed Functional, Integration, System, Regression, and Stress testing on a cloud-based Property Management System.",
    "Developed 200+ manual test cases covering the complete PMS lifecycle — from front desk operations to reporting services.",
    "Executed API testing using Postman to validate web service endpoints and data integrity.",
    "Executed SQL queries to validate data accuracy and business logic in the backend database.",
]
for b in project1_bullets:
    add_bullet(doc, b)

# Project 2
add_project_heading(doc,
    "Touche Hospitality POS System (Multi-Platform)",
    "Role: Lead Functional Test Analyst  |  Tools: SQL Server, TFS, Postman, JMeter"
)

project2_bullets = [
    "Led end-to-end testing for a multi-platform POS system spanning Android mobile, web/cloud, and Windows desktop.",
    "Developed 800+ manual test cases covering all platform-specific and cross-platform scenarios.",
    "Performed API testing using Postman and performance testing using Apache JMeter.",
    "Coordinated with third-party developers for integration testing with OPERA PMS, Shawman, and Urban Piper.",
    "Provided on-demand support to the support team for urgent production issues and application deployment.",
]
for b in project2_bullets:
    add_bullet(doc, b)

# Project 3
add_project_heading(doc,
    "WebWish PMS — Test Automation Framework & AI Integration",
    "Role: Test Automation Lead  |  Tools: Playwright, TypeScript, GitHub Copilot, MCP Server, Docker, Jenkins"
)

project3_bullets = [
    "Architected and built a production-grade Playwright test automation framework using TypeScript with Page Object Model, custom fixtures, and modular folder structure.",
    "Developed 50+ reusable page objects and custom utility libraries (ElementActions, ExcelUtils, DatabaseHelper) for cross-module test support.",
    "Implemented multi-browser support (Chromium, Firefox, WebKit) with parallel execution, reducing full regression run time by 60%.",
    "Integrated Allure and custom HTML reporters for comprehensive test reporting and trend analysis.",
    "Set up Docker-based test execution and Jenkins CI/CD pipelines for automated nightly regression runs.",
    "Leveraged GitHub Copilot and Playwright MCP Server for AI-assisted test generation, achieving 40% faster script authoring.",
    "Built Agentic AI-driven workflows to auto-generate test cases from user stories and automatically categorize defects by severity.",
    "Maintained and evolved the framework across 8+ application modules (Front Desk, Marketing, Reports, Global Search, Night Audit, Housekeeping, Cashiering, System Config).",
]
for b in project3_bullets:
    add_bullet(doc, b)

# ══════════════════════════════════════════════════════════════
#  TECHNICAL SKILLS (TABLE)
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Technical Skills")

skills_data = [
    ("Testing Types", "Manual Testing, Functional, Regression, Integration, System, Stress, Ad-hoc, Exploratory, API, Performance"),
    ("Testing Techniques", "Black Box, Boundary Value Analysis (BVA), Equivalence Class Partitioning (ECP), Decision Tables, Use Case Testing"),
    ("Automation & Framework", "Playwright (TypeScript), Selenium, Page Object Model (POM), Custom Fixtures, Modular Architecture"),
    ("AI-Assisted Testing", "GitHub Copilot, Playwright MCP Server, Agentic AI Workflows, AI Code Generation, LLM-based Test Planning"),
    ("API & Performance", "Postman, Apache JMeter"),
    ("Database", "SQL Server, Database Validation & Query Writing"),
    ("CI/CD & DevOps", "Jenkins, Docker, Git/GitHub, CI/CD Pipeline Integration"),
    ("Reporting", "Allure Reports, Custom HTML Reporters, Test Dashboards"),
    ("Defect Tracking", "TFS (Team Foundation Server), Jira"),
    ("Methodologies", "Agile, Scrum, SDLC, STLC"),
    ("Other Tools", "VS Code, GitHub Copilot, TestProject, Excel (Data-Driven Testing)"),
]

table = doc.add_table(rows=len(skills_data), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Set table style
table.style = 'Table Grid'

for i, (category, skills) in enumerate(skills_data):
    row = table.rows[i]
    
    # Category cell
    cell_cat = row.cells[0]
    cell_cat.width = Cm(4.5)
    p_cat = cell_cat.paragraphs[0]
    p_cat.paragraph_format.space_before = Pt(3)
    p_cat.paragraph_format.space_after = Pt(3)
    run_cat = p_cat.add_run(category)
    run_cat.font.name = 'Calibri'
    run_cat.font.size = Pt(10)
    run_cat.font.bold = True
    run_cat.font.color.rgb = DARK_BLUE
    
    # Skills cell
    cell_skills = row.cells[1]
    cell_skills.width = Cm(12)
    p_skills = cell_skills.paragraphs[0]
    p_skills.paragraph_format.space_before = Pt(3)
    p_skills.paragraph_format.space_after = Pt(3)
    run_skills = p_skills.add_run(skills)
    run_skills.font.name = 'Calibri'
    run_skills.font.size = Pt(10)
    run_skills.font.color.rgb = DARK_GRAY
    
    # Alternate row shading
    if i % 2 == 0:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
        row.cells[0]._tc.get_or_add_tcPr().append(shading)
        row.cells[1]._tc.get_or_add_tcPr().append(shading)

# ══════════════════════════════════════════════════════════════
#  PROFESSIONAL DEVELOPMENT
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Professional Development")

pd_bullets = [
    ("Agentic AI for Test Automation (2026) ", "— Applied Agentic AI concepts to automate repetitive QA workflows, including test case generation from requirements, automated regression suite creation, and intelligent defect categorization using LLM-based tools."),
    ("Playwright MCP Server Integration (2026) ", "— Integrated Model Context Protocol (MCP) Server with Playwright for AI-driven browser automation, enabling intelligent test scenario generation and autonomous test execution."),
    ("GitHub Copilot for Test Development (2026) ", "— Leveraged AI-assisted code generation to accelerate test script authoring by 40%, including page object creation, assertion building, and test data generation."),
    ("Playwright Test Framework Development (2026) ", "— Designed and built a scalable, production-grade test automation framework using TypeScript, POM pattern, custom fixtures, Allure reporting, Docker, and CI/CD integration."),
    ("Apache JMeter (2023) ", "— Performance and load testing for web applications."),
    ("Self-taught Postman for API Testing ", "— Independently learned and implemented Postman across the QA team; recognized by Project Manager."),
]

for bold_part, rest_part in pd_bullets:
    add_bullet(doc, rest_part, bold_prefix=bold_part)

# ══════════════════════════════════════════════════════════════
#  EDUCATION
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Education")

add_subsection_heading(doc,
    "Bachelor of Engineering — Electronics & Communication Engineering",
    "Rustamji Institute of Technology & Management, BSF Academy, Gwalior, MP | 2013 – 2017"
)

add_subsection_heading(doc,
    "Higher Secondary (12th — PCM)",
    "BSF School, Tekanpur, Gwalior, MP | 2012 – 2013"
)

# ══════════════════════════════════════════════════════════════
#  ADDITIONAL INFORMATION
# ══════════════════════════════════════════════════════════════

add_section_heading(doc, "Additional Information")

add_bullet(doc, "English, Hindi", bold_prefix="Languages: ")
add_bullet(doc, "Agile Foundations (2025)", bold_prefix="Certifications: ")

# ─── Save ───
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Avinash_Tiwary_Resume.docx")
doc.save(output_path)
print(f"✅ Resume saved to: {output_path}")
